"""Report exports: CSV, XLSX (openpyxl) and PDF (WeasyPrint, with a print-page fallback)."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from django.utils.text import slugify

from apps.core.pdf import weasyprint_html

FORMATS = ("csv", "xlsx", "pdf")


def format_cell(value, kind: str) -> str:
    """Display text for one cell (screen and PDF)."""
    if value is None or value == "":
        return "—" if kind != "text" else ""
    if kind == "money":
        return f"{Decimal(value):,.2f}"
    if kind == "int":
        return f"{int(value):,}"
    if kind == "percent":
        return f"{value}%"
    if kind in ("hours", "days", "decimal"):
        return f"{float(value):,.1f}"
    if kind == "datetime" and isinstance(value, datetime):
        return f"{timezone.localtime(value):%d %b %Y %H:%M}"
    if kind == "date" and isinstance(value, date):
        return f"{value:%d %b %Y}"
    return str(value)


def _raw(value, kind):
    """Machine-friendly value for CSV/XLSX: numbers stay numbers, dates ISO / real dates."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None)
    if kind == "money" and not isinstance(value, str):
        return Decimal(value).quantize(Decimal("0.01"))
    return value


def filename(rdef, filters, ext) -> str:
    return f"{slugify(rdef.title)}-{filters.start:%Y%m%d}-{filters.end:%Y%m%d}.{ext}"


def _table(result):
    rows = [[_raw(r.get(c.key), c.kind) for c in result.columns] for r in result.rows]
    totals = (
        [_raw(result.totals.get(c.key), c.kind) for c in result.columns] if result.totals else None
    )
    return rows, totals


def to_csv(rdef, filters, result) -> HttpResponse:
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename(rdef, filters, "csv")}"'
    response.write("﻿")  # BOM so Excel opens UTF-8 (Urdu names, en dashes) correctly
    writer = csv.writer(response)
    writer.writerow([c.label for c in result.columns])
    rows, totals = _table(result)
    for row in rows:
        writer.writerow([f"{v:%Y-%m-%d %H:%M}" if isinstance(v, datetime) else v for v in row])
    if totals:
        writer.writerow(totals)
    return response


def to_xlsx(rdef, filters, result, organization) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = rdef.title[:31]
    ws.append([f"{organization.name} — {rdef.title}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([filters.describe(rdef)])
    ws.append([])
    ws.append([c.label for c in result.columns])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    rows, totals = _table(result)
    for row in rows:
        ws.append(row)
    if totals:
        ws.append(totals)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    formats = {
        "money": "#,##0.00",
        "int": "#,##0",
        "date": "dd mmm yyyy",
        "datetime": "dd mmm yyyy hh:mm",
        "hours": "0.0",
        "days": "0.0",
    }
    for idx, col in enumerate(result.columns, start=1):
        letter = get_column_letter(idx)
        width = max(
            [len(col.label)]
            + [len(format_cell(r.get(col.key), col.kind)) for r in result.rows[:500]]
        )
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 50)
        if col.kind in formats:
            for cell in ws[letter][header_row:]:
                cell.number_format = formats[col.kind]
    if result.note:
        ws.append([])
        ws.append([result.note])
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename(rdef, filters, "xlsx")}"'
    return response


def to_pdf(request, rdef, filters, result, context) -> HttpResponse:
    """A4 landscape PDF. Without WeasyPrint's native libraries, return the print page instead."""
    html = render_to_string("reports/print.html", {**context, "pdf": True}, request=request)
    HTML = weasyprint_html()
    if HTML is None:
        page = render_to_string(
            "reports/print.html", {**context, "pdf_unavailable": True}, request=request
        )
        return HttpResponse(page)
    pdf = HTML(string=html, base_url=request.build_absolute_uri("/")).write_pdf()
    response = HttpResponse(pdf, content_type="application/pdf")
    response["Content-Disposition"] = f'inline; filename="{filename(rdef, filters, "pdf")}"'
    return response
