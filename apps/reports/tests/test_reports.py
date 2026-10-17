import io
from datetime import time
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.accounts.models import RoleAssignment
from apps.accounts.services import effective_permissions
from apps.appointments.models import Appointment
from apps.audit.models import AuditLog
from apps.billing import services as billing
from apps.billing.models import PaymentMethod
from apps.encounters import services as encounters
from apps.encounters.models import EncounterDiagnosis
from apps.reports.engine import REPORTS

D = Decimal


def _req(rf, user, org, branch):
    req = rf.post("/")
    req.user, req.organization, req.branch = user, org, branch
    req.sd_perms = effective_permissions(user, branch)
    return req


@pytest.fixture
def cash(org):
    billing.ensure_payment_methods(org)
    return PaymentMethod.objects.get(organization=org, code="cash")


def _paid_invoice(req, patient, cash, amount="500", appointment=None):
    inv = billing.create_invoice(request=req, patient=patient, appointment=appointment)
    billing.add_item(inv, request=req, description="Consultation", unit_price=D(amount))
    inv.refresh_from_db()
    billing.receive_payment(inv, request=req, method=cash, amount=amount)
    inv.refresh_from_db()
    return inv


def _signed_visit(rf, doctor, org, branch, patient, diagnosis="Typhoid fever"):
    req = _req(rf, doctor.user, org, branch)
    enc, _ = encounters.start(request=req, patient=patient)
    enc = encounters.save_note(
        enc,
        request=req,
        data={"chief_complaint": "Fever", "provisional_diagnosis": diagnosis},
        revision=0,
    )
    EncounterDiagnosis.objects.create(encounter=enc, text=diagnosis, added_by=doctor.user)
    encounters.complete(enc, request=req)
    return enc


# --- Every report renders and exports -----------------------------------------------------------


@pytest.mark.parametrize("key", sorted(REPORTS))
def test_every_report_renders_empty(admin_client, branch, key):
    response = admin_client.get(f"/reports/{key}/")
    assert response.status_code == 200
    assert REPORTS[key].title.encode() in response.content


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "pdf"])
def test_exports_and_audit(admin_client, branch, fmt):
    response = admin_client.get(f"/reports/opd-visits/?format={fmt}")
    assert response.status_code == 200
    expected = {
        "csv": "text/csv",
        "xlsx": "spreadsheetml",
        "pdf": "application/pdf",  # or the print page when WeasyPrint's libraries are missing
    }[fmt]
    content_type = response["Content-Type"]
    assert expected in content_type or (fmt == "pdf" and "text/html" in content_type)
    log = AuditLog.objects.get(action="report.export")
    assert log.summary.startswith(f"opd-visits {fmt}")


def test_hub_lists_reports_by_permission(client, make_user, branch):
    client.force_login(make_user("cashier"))  # billing + reports.financial.view
    content = client.get("/reports/").content
    assert b"Daily collection" in content and b"Refunds" in content
    assert b"OPD visits" not in content  # needs reports.view

    client.force_login(make_user("reception"))
    assert client.get("/reports/").status_code == 403
    assert client.get("/reports/daily-collection/").status_code == 403


def test_pharmacist_sees_stock_reports_only(client, make_user, branch):
    client.force_login(make_user("pharmacy"))
    assert client.get("/reports/current-stock/").status_code == 200
    assert client.get("/reports/purchases/").status_code == 200
    assert client.get("/reports/pharmacy-sales/").status_code == 403
    assert client.get("/reports/daily-collection/").status_code == 403


def test_unknown_report_is_404(admin_client, branch):
    assert admin_client.get("/reports/not-a-report/").status_code == 404


# --- Branch scoping -----------------------------------------------------------------------------


def test_branch_filter_never_widens_access(client, make_user, org, branch, branch2):
    user = make_user("hospital_admin", branch=branch)
    client.force_login(user)
    assert client.get(f"/reports/daily-collection/?branch={branch2.pk}").status_code == 403
    # "All branches" means all branches this user may report on: only the main one here.
    response = client.get("/reports/daily-collection/?branch=all")
    assert response.status_code == 200
    assert response.context["f"].branches == [branch]


def test_branch_needs_report_permission_there(client, make_user, org, branch, branch2):
    user = make_user("hospital_admin", branch=branch)
    RoleAssignment.objects.create(
        user=user, role=org.roles.get(system_key="reception"), branch=branch2
    )
    client.force_login(user)
    # The user works in branch2, but only as reception: no reports there.
    assert client.get(f"/reports/opd-visits/?branch={branch2.pk}").status_code == 403


def test_all_branches_adds_up(rf, admin_client, admin_user, org, branch, branch2, patient, cash):
    _paid_invoice(_req(rf, admin_user, org, branch), patient, cash, "500")
    _paid_invoice(_req(rf, admin_user, org, branch2), patient, cash, "700")
    one = admin_client.get("/reports/daily-collection/").context["result"]
    both = admin_client.get("/reports/daily-collection/?branch=all").context["result"]
    assert one.totals["net"] == D("500")
    assert both.totals["net"] == D("1200")


# --- Numbers -----------------------------------------------------------------------------------


def test_daily_collection_matches_day_totals(rf, admin_client, admin_user, org, branch, patient,
                                             cash):  # fmt: skip
    req = _req(rf, admin_user, org, branch)
    inv = _paid_invoice(req, patient, cash, "1500")
    billing.refund(inv.payments.first(), request=req, amount="200", reason="Overcharged")
    result = admin_client.get("/reports/daily-collection/").context["result"]
    totals = billing.day_totals(branch)
    assert result.totals["collected"] == totals["collected"] == D("1500")
    assert result.totals["net"] == totals["net"] == D("1300")
    methods = admin_client.get("/reports/payment-methods/").context["result"]
    assert methods.rows[0]["method"] == "Cash" and methods.rows[0]["net"] == D("1300")
    refunds = admin_client.get("/reports/refunds/").context["result"]
    assert refunds.totals["amount"] == D("200") and refunds.rows[0]["reason"] == "Overcharged"
    csv_text = admin_client.get("/reports/daily-collection/?format=csv").content.decode()
    assert "1300.00" in csv_text


def test_outstanding_and_revenue_by_department(
    rf, admin_client, admin_user, org, branch, patient, doctor
):
    req = _req(rf, admin_user, org, branch)
    appt = Appointment.objects.create(
        organization=org, branch=branch, number="APT-1", patient=patient, doctor=doctor,
        department=doctor.department, date=timezone.localdate(), start_time=time(9),
        end_time=time(9, 30), fee=D("1500"), created_by=admin_user,
    )  # fmt: skip
    inv = billing.create_invoice(request=req, patient=patient, appointment=appt)
    inv.refresh_from_db()
    billing.issue(inv, request=req)
    rows = admin_client.get("/reports/outstanding/").context["result"].rows
    assert [(r["number"], r["balance"]) for r in rows] == [(inv.number, D("1500"))]
    revenue = admin_client.get("/reports/revenue-by-department/").context["result"]
    assert revenue.rows[0]["name"] == "General medicine"
    assert revenue.rows[0]["net"] == D("1500")
    by_doctor = admin_client.get(f"/reports/revenue-by-doctor/?doctor={doctor.pk}")
    assert by_doctor.context["result"].rows[0]["name"] == "Dr. Zain Ali"


def test_clinical_reports_count_visits_and_diagnoses(
    rf, admin_client, org, branch, doctor, patient
):
    _signed_visit(rf, doctor, org, branch, patient)
    opd = admin_client.get("/reports/opd-visits/").context["result"]
    assert opd.totals["visits"] == 1 and opd.totals["completed"] == 1
    docs = admin_client.get("/reports/doctor-visits/").context["result"]
    assert docs.rows[0]["name"] == "Dr. Zain Ali" and docs.rows[0]["visits"] == 1
    depts = admin_client.get(f"/reports/department-visits/?department={doctor.department.pk}")
    assert depts.context["result"].rows[0]["name"] == "General medicine"
    dx = admin_client.get("/reports/diagnosis-count/").context["result"]
    assert dx.rows == [{"code": "", "name": "Typhoid fever", "count": 1, "patients": 1}]
    new = admin_client.get("/reports/new-patients/").context["result"]
    assert new.totals["total"] == 1 and new.totals["female"] == 1


def test_xlsx_has_title_header_and_rows(rf, admin_client, org, branch, doctor, patient):
    from openpyxl import load_workbook

    _signed_visit(rf, doctor, org, branch, patient)
    response = admin_client.get("/reports/doctor-visits/?format=xlsx")
    ws = load_workbook(io.BytesIO(response.content)).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0][0].endswith("Doctor-wise visits")
    assert values[3][:2] == ["Doctor", "Visits"]
    assert values[4][:2] == ["Dr. Zain Ali", 1]


def test_long_ranges_are_capped(admin_client, branch):
    response = admin_client.get("/reports/opd-visits/?from=2020-01-01&to=2026-01-01")
    f = response.context["f"]
    assert (f.end - f.start).days == 365


# --- Dashboards ---------------------------------------------------------------------------------


def test_admin_kpis_on_dashboard(rf, client, make_user, org, branch, patient, cash, doctor):
    _paid_invoice(_req(rf, make_user("super_admin"), org, branch), patient, cash, "800")
    _signed_visit(rf, doctor, org, branch, patient)
    client.force_login(make_user("hospital_admin"))
    response = client.get("/")
    kpi = response.context["kpi"]
    assert kpi["opd_today"] == 1 and kpi["patients_today"] == 1
    assert kpi["revenue_today"] == D("800")
    assert len(kpi["trend"]) == 14 and kpi["trend"][-1]["visits"] == 1
    assert b"Today at" in response.content

    client.force_login(make_user("doctor"))
    assert "kpi" not in client.get("/").context


def test_lab_and_pharmacy_dashboards_permissions(client, make_user, branch):
    client.force_login(make_user("lab"))
    assert client.get("/dashboards/lab/").status_code == 200
    assert client.get("/dashboards/pharmacy/").status_code == 403
    client.force_login(make_user("pharmacy"))
    response = client.get("/dashboards/pharmacy/")
    assert response.status_code == 200
    assert response.context["k"]["sales_count"] == 0
    client.force_login(make_user("reception"))
    assert client.get("/dashboards/lab/").status_code == 403


# --- Lab, pharmacy and IPD reports with data ------------------------------------------------------


def test_lab_reports_with_data(rf, admin_client, make_user, org, branch, patient):
    from apps.lab import services as lab
    from apps.lab.models import LabTest, LabTestCategory

    cat = LabTestCategory.objects.create(organization=org, name="Haematology")
    hb = LabTest.objects.create(
        organization=org, code="HB", name="Haemoglobin", category=cat, sample_type="blood",
        price=D("300"), unit="g/dL", turnaround_hours=4,
    )  # fmt: skip
    req = _req(rf, make_user("lab"), org, branch)
    order = lab.create_order(request=req, patient=patient, tests=[hb], priority="stat")
    pending = admin_client.get("/reports/lab-pending/").context["result"]
    assert [r["test"] for r in pending.rows] == ["Haemoglobin"]
    lab.collect_sample(order.samples.get(), request=req)
    item = order.items.get()
    lab.enter_result(item, request=req, value="13.1")
    lab.verify([item], request=req)
    volume = admin_client.get("/reports/lab-test-volume/").context["result"]
    assert volume.rows[0]["ordered"] == 1 and volume.rows[0]["verified"] == 1
    tat = admin_client.get("/reports/lab-tat/").context["result"]
    assert tat.rows[0]["count"] == 1 and tat.rows[0]["within"] == 100.0
    revenue = admin_client.get("/reports/lab-revenue/").context["result"]
    assert revenue.totals["net"] == D("300")
    assert admin_client.get("/reports/lab-pending/").context["result"].rows == []
    dash = admin_client.get("/dashboards/lab/").context["k"]
    assert dash["verified_today"] == 1 and dash["pending_samples"] == 0


def test_pharmacy_reports_with_data(rf, admin_client, admin_user, org, branch, cash):
    from datetime import timedelta

    from apps.pharmacy import pos, stock
    from apps.pharmacy.models import Product

    req = _req(rf, admin_user, org, branch)
    panadol = Product.objects.create(
        organization=org, sku="P-1", brand_name="Panadol", strength="500 mg",
        sale_price=D("3"), reorder_level=50,
    )  # fmt: skip
    stock.receive(
        request=req, product=panadol, batch_number="B1",
        expiry_date=timezone.localdate() + timedelta(days=30), quantity=40, purchase_cost=D("1"),
    )  # fmt: skip
    cart = pos.open_cart(req)
    pos.add_product(cart, request=req, product=panadol, quantity=10)
    pos.checkout(cart, request=req, payments=[{"method": cash, "amount": "30"}])

    sales = admin_client.get("/reports/pharmacy-sales/").context["result"]
    assert sales.totals["total"] == D("30")
    gross = admin_client.get("/reports/pharmacy-gross/").context["result"]
    assert (gross.totals["net"], gross.totals["cost"], gross.totals["margin"]) == (
        D("30"), D("10"), D("20"),
    )  # fmt: skip
    movement = admin_client.get("/reports/product-movement/").context["result"].rows[0]
    assert (movement["received"], movement["sold"], movement["net"]) == (40, 10, 30)
    current = admin_client.get("/reports/current-stock/").context["result"].rows[0]
    assert (current["on_hand"], current["cost"]) == (30, D("30"))
    low = admin_client.get("/reports/low-stock/").context["result"].rows
    assert low[0]["status"] == "Low" and low[0]["short"] == 20
    expiry = admin_client.get("/reports/expiry/").context["result"].rows
    assert expiry[0]["status"] == "Near expiry" and expiry[0]["qty"] == 30
    dash = admin_client.get("/dashboards/pharmacy/").context["k"]
    assert dash["sales_count"] == 1 and dash["low_stock"] == 1 and dash["near_expiry"] == 1
    assert dash["top_products"][0]["units"] == 10


def test_ipd_reports_with_data(rf, admin_client, admin_user, org, branch, patient, doctor):
    from apps.ipd.tests.helpers import admit, make_ward

    req = _req(rf, admin_user, org, branch)
    req.META["REMOTE_ADDR"] = "127.0.0.1"
    _, _, beds = make_ward(org, branch, beds=2)
    admit(req, patient, doctor, beds[0])
    adm = admin_client.get("/reports/ipd-admissions/").context["result"]
    assert adm.rows[0]["bed"] == beds[0].code and adm.rows[0]["status"] == "Admitted"
    occ = admin_client.get("/reports/ipd-occupancy/").context["result"]
    assert occ.totals["beds"] == 2 and occ.totals["occupied"] == 1 and occ.totals["now"] == 50.0
    assert admin_client.get("/reports/ipd-discharges/").context["result"].rows == []
    assert admin_client.get("/").context["kpi"]["occupancy"] == 50

    # A two-day stay ending now (set directly; the discharge flow has its own tests).
    from datetime import timedelta

    from apps.ipd.models import Admission

    now = timezone.now()
    Admission.objects.update(
        status=Admission.Status.DISCHARGED,
        discharge_type="routine",
        admitted_at=now - timedelta(days=2),
        discharged_at=now,
    )
    dis = admin_client.get("/reports/ipd-discharges/").context["result"].rows
    assert dis[0]["stay"] == 2.0 and dis[0]["type"] == "Routine (advised)"
    stay = admin_client.get("/reports/ipd-average-stay/").context["result"]
    assert stay.rows[0]["avg"] == 2.0 and stay.totals["count"] == 1
